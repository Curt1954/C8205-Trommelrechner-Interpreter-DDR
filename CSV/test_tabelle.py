#!/usr/bin/env python3
"""Erzeugt eine Übersichtstabelle aller Tests.

Spalten:
  1. Testdateiname
  2. Inhalt des Tests (verbal, aus den #DOK-Zeilen)
  3. Dateien der virtuellen Geräte (Leser, SW-Ausgabe, K0, ...) und
     Test-Vergleichsdateien (--pruefe-sw-ausgabe / --pruefe-stanzer1)
  4. zugehöriger ASM-Quelltext (falls vorhanden)
  5. ASM-Zuordnung: leer = ASM-Datei mit gleichem Namen; sonst die Art der
     Zuordnung (über Trommelinhalt, über Namen, Voraussetzung: Minimalprogramm)

Spalte 2 enthaelt standardmaessig nur die erste #DOK-Zeile (hoechstens
--max-inhalt Zeichen). Haben mehrere Tests denselben Text, werden bei ihnen
weitere #DOK-Zeilen angehaengt, bis sie sich unterscheiden. --voll liefert
wie frueher alle #DOK-Zeilen ungekuerzt.

Die MD-Ausgabe ist eine Projektdatei nach C8205-Versionierung.md: Titelzeile, darunter
'Version: y.xx-zzz'. Neue Datei: 0.01-001; aendert sich der Inhalt gegenueber der vorhandenen
Ausgabedatei, steigt der Build um 1 (Unter- und Hauptversion setzt nur der Projektinhaber).

Aufruf (aus beliebigem Verzeichnis; die Basis mit Tests/ und Assembler/ wird
automatisch gesucht: aktuelles Verzeichnis, Skriptverzeichnis, jeweils deren
Elternverzeichnis - z.B. wenn das Skript in CSV/ neben Assembler/ liegt):
    python3 test_tabelle.py                    # schreibt test_tabelle.csv, .md, .xlsx (wenn openpyxl da ist)
    python3 test_tabelle.py --basis /pfad --format md csv xlsx html [--max-inhalt 140 | --voll]
Die Ausgabedateien landen im Skriptverzeichnis (oder unter --ausgabe, wenn Pfad angegeben).
"""
import argparse
import csv
import difflib
import html
import os
import re
import sys
from collections import Counter, defaultdict

ASM_VERZEICHNISSE = ["Claude", "Claude_bsp_reassemblierung", "Proj_FeldSortieren",
                     "Projekt_1", "Projekt_2"]

# Schalter in den Testdateien -> Anzeigename
GERAETE = [
    ("--leser1", "Leser 1"),
    ("--leser2", "Leser 2"),
    ("--k0", "K0-Eingabe"),
    ("--sw-ausgabe", "SW-Ausgabe"),
    ("--pruefe-sw-ausgabe", "Vergleich SW-Ausgabe"),
    ("--pruefe-stanzer1", "Vergleich Stanzer 1"),
]


SKRIPTDIR = os.path.dirname(os.path.abspath(__file__))


def finde_basis(vorgabe=None):
    """Verzeichnis suchen, das Tests/ und Assembler/ enthaelt."""
    if vorgabe:
        return os.path.abspath(vorgabe)
    kandidaten = []
    for start in (os.getcwd(), SKRIPTDIR):
        k = os.path.abspath(start)
        for _ in range(3):  # Verzeichnis selbst und bis zu 2 Ebenen darueber
            if k not in kandidaten:
                kandidaten.append(k)
            k = os.path.dirname(k)
    for k in kandidaten:
        if os.path.isdir(os.path.join(k, "Tests")) and os.path.isdir(os.path.join(k, "Assembler")):
            return k
    for k in kandidaten:
        if os.path.isdir(os.path.join(k, "Tests")):
            return k
    sys.exit("Keine Basis mit Tests/ und Assembler/ gefunden. Bitte --basis /pfad angeben.")


def norm(s):
    return re.sub(r"[^a-z0-9]+", "_", s.lower()).strip("_")


def lies(pfad):
    with open(pfad, encoding="utf-8", errors="replace") as f:
        return f.read().splitlines()


def trommelzellen(pfad):
    """Zellen (Adresse -> Wert) aus dem Abschnitt 'Trommeldatei:'."""
    zellen, im_abschnitt = {}, False
    for z in lies(pfad):
        if z.startswith("Trommeldatei"):
            im_abschnitt = True
            continue
        if im_abschnitt and re.match(r"(Tastenfolge|Pr\S*fungen)\s*:", z):
            break
        if im_abschnitt:
            m = re.match(r"\s*([0-7]{4}):\s*([0-7]+)", z)
            if m:
                zellen[m[1]] = m[2]
    return zellen


def dok_texte(zeilen):
    """Texte der #DOK-Zeilen in Dateireihenfolge, ohne Duplikate und ohne leere."""
    teile = []
    for z in zeilen:
        m = re.match(r"#DOK-\d+\s*(.*)", z)
        if m and m[1].strip() and m[1].strip() not in teile:
            teile.append(m[1].strip())
    return teile


def kuerze(text, maxlen):
    """Text auf hoechstens ca. maxlen Zeichen kuerzen (an einer Wortgrenze, mit ' …')."""
    if not maxlen or len(text) <= maxlen:
        return text
    schnitt = text[:maxlen]
    if " " in schnitt:
        schnitt = schnitt.rsplit(" ", 1)[0]
    return schnitt.rstrip(" ,;:(") + " …"


def kurzbeschreibungen(namen, dok_je_test, ersatz, maxlen):
    """Name -> Kurzbeschreibung. Grundlage ist die erste #DOK-Zeile (ohne #DOK-Zeile
    der Ersatztext). Tests mit gleichem Text bekommen weitere #DOK-Zeilen angehaengt,
    bis sie sich unterscheiden oder keine mehr da sind; danach wird gekuerzt."""
    n_zeilen = {n: 1 for n in namen}

    def text(n):
        d = dok_je_test[n]
        return " ".join(d[:n_zeilen[n]]) if d else ersatz[n]

    while True:
        zaehler = Counter(text(n) for n in namen)
        geaendert = False
        for n in namen:
            if zaehler[text(n)] > 1 and n_zeilen[n] < len(dok_je_test[n]):
                n_zeilen[n] += 1
                geaendert = True
        if not geaendert:
            break
    return {n: kuerze(text(n), maxlen) for n in namen}


def ersatzbeschreibung(zeilen, name):
    """Falls keine #DOK-Zeile existiert: Kommentare / Tastenfolge als Hinweis."""
    kommentare = []
    for z in zeilen:
        if z.startswith("#") and not re.match(r"#(DOK|SPEI|PROT|SPUR|STARTE|INTERRUPT|AC|BZ)", z):
            t = z.lstrip("# ").strip()
            if t:
                kommentare.append(t)
    text = " ".join(kommentare[:2])
    for i, z in enumerate(zeilen):
        if z.startswith("Tastenfolge"):
            for folge in zeilen[i + 1:i + 2]:
                if folge.strip():
                    text += f" [Tastenfolge: {folge.strip()}]"
            break
    return (text.strip() or "(keine #DOK-Beschreibung vorhanden)")


def geraete_zeilen(zeilen, basis, testname=""):
    """Geräte-/Vergleichsdateien aus den '--schalter wert'-Zeilen (nur Pfade)."""
    ausgabe = []
    schalter_bekannt = {sw for sw, _ in GERAETE}
    for z in zeilen:
        if not z.startswith("--"):
            continue
        teile = z.split(" #")[0].strip().split(None, 1)
        if len(teile) < 2 or teile[0] not in schalter_bekannt:
            continue
        schalter, wert = teile[0], re.sub(r"^datei=", "", teile[1].strip())
        pfad = os.path.join(basis, wert)
        # Ausgabedateien/-verzeichnisse entstehen erst beim Lauf
        if schalter != "--sw-ausgabe" and not os.path.exists(pfad):
            print(f"Warnung: {testname}: Datei fehlt: {wert}", file=sys.stderr)
        ausgabe.append(wert)
    return ausgabe


def asm_dateien(basis):
    liste = []
    for d in ASM_VERZEICHNISSE:
        wurzel = os.path.join(basis, "Assembler", d)
        for r, _, fs in os.walk(wurzel):
            for f in sorted(fs):
                if f.endswith(".asm"):
                    pfad = os.path.join(r, f)
                    liste.append((os.path.relpath(pfad, os.path.join(basis, "Assembler")), pfad))
    return sorted(liste)


def kern(s):
    """Normalisierter Name ohne 'bsp'-Teile (pj2-bsp_x und pj2_x sollen passen)."""
    return "_".join(t for t in norm(s).split("_") if t != "bsp")


def finde_asm(testpfad, stamm, asm, asm_zellen, text=""):
    """Zuordnung: 1. gleicher Name, 2. Trommelinhalt passt (bei Mehrdeutigkeit
    ähnlichster Name), 3. Namenspräfix, 4. Voraussetzung Minimalprogramm."""
    def stamm_von(rel):
        return os.path.splitext(os.path.basename(rel))[0]

    gleich = [rel for rel, p in asm if stamm_von(p) == stamm]
    if gleich:
        return gleich[0], ""
    zellen = trommelzellen(testpfad)
    def aehnl(rel):
        return difflib.SequenceMatcher(None, norm(stamm), norm(stamm_von(rel))).ratio()
    if zellen:
        kand = [rel for rel, p in asm
                if asm_zellen[rel] and all(asm_zellen[rel].get(k) == v for k, v in zellen.items())]
        kand.sort(key=aehnl, reverse=True)
        if kand and aehnl(kand[0]) >= 0.6:
            hinweis = " (über Trommelinhalt)"
            if len(kand) > 1 and aehnl(kand[0]) == aehnl(kand[1]):
                hinweis = " (mehrdeutig: " + ", ".join(os.path.basename(k) for k in kand[:3]) + ")"
            return kand[0], hinweis
    # Namenspräfix (ziffern_uebergehen -> ziffern, pj2-bsp_maximum_3_zahlen_a_max -> pj2_maximum_3_zahlen)
    k = kern(stamm)
    pre = [rel for rel, p in asm
           if k.startswith(kern(stamm_von(rel)) + "_") and len(kern(stamm_von(rel))) >= 5]
    if pre:
        pre.sort(key=lambda r: len(kern(stamm_von(r))), reverse=True)
        return pre[0], " (über Namen)"
    # BIS-Tests setzen das Minimalprogramm auf der Trommel voraus
    if "minimalprogramm" in text.lower() or stamm.startswith("pj2-bis_"):
        for rel, p in asm:
            if stamm_von(p) == "pj2-bis_minimalprogramm":
                return rel, " (Voraussetzung: Minimalprogramm auf der Trommel)"
    return "", ""


def sammle(basis, voll=False, maxlen=140):
    testdir = os.path.join(basis, "Tests")
    if not os.path.isdir(testdir):
        sys.exit(f"Verzeichnis nicht gefunden: {testdir}")
    asm = asm_dateien(basis)
    asm_zellen = {}
    for rel, p in asm:
        t = os.path.splitext(p)[0] + ".test"
        asm_zellen[rel] = trommelzellen(t) if os.path.exists(t) else {}

    zeilen_tab = []
    dok_je_test, ersatz = {}, {}
    for name in sorted(os.listdir(testdir)):
        if not name.endswith((".txt", ".test")):
            continue
        pfad = os.path.join(testdir, name)
        z = lies(pfad)
        # Datei hat das Testformat? (Soll-Dateien haben keinen Abschnitt 'Trommeldatei:')
        if not any(x.startswith("Trommeldatei") for x in z):
            continue
        stamm = os.path.splitext(name)[0]
        dok = dok_texte(z)
        dok_je_test[name] = dok
        ersatz[name] = "" if dok else ersatzbeschreibung(z, name)
        text = " ".join(dok) or ersatz[name]  # voller Text; fuer die ASM-Suche
        geraete = geraete_zeilen(z, basis, name)
        rel, hinweis = finde_asm(pfad, stamm, asm, asm_zellen, text)
        zeilen_tab.append([name, text, "\n".join(geraete), rel, hinweis.strip(" ()")])
        if hinweis:
            print(f"Hinweis: {name} -> {rel}{hinweis}", file=sys.stderr)
    if not voll:
        kurz = kurzbeschreibungen([r[0] for r in zeilen_tab], dok_je_test, ersatz, maxlen)
        for r in zeilen_tab:
            r[1] = kurz[r[0]]
    return [tuple(r) for r in zeilen_tab]


PLATZHALTER = ("[Tastenfolge", "(keine #DOK")


def berichte_luecken(tab):
    """Hinweise (stderr): Tests ohne echte Beschreibung und Tests mit gleicher Beschreibung."""
    platz = [r[0] for r in tab if r[1].startswith(PLATZHALTER)]
    if platz:
        print(f"Hinweis: {len(platz)} Tests ohne #DOK-Text (nur Platzhalter): " + ", ".join(platz),
              file=sys.stderr)
    gruppen = defaultdict(list)
    for r in tab:
        if not r[1].startswith(PLATZHALTER):
            gruppen[r[1]].append(r[0])
    gleich = [ns for ns in gruppen.values() if len(ns) > 1]
    if gleich:
        print(f"Hinweis: {len(gleich)} Gruppen mit gleicher Beschreibung "
              f"({sum(len(ns) for ns in gleich)} Tests) - #DOK-Zeilen unterscheidbar machen:", file=sys.stderr)
        for ns in gleich:
            print("  " + ", ".join(ns), file=sys.stderr)


KOPF = ["Testdatei", "Inhalt des Tests", "Dateien virtueller Geräte / Vergleichsdatei", "ASM-Quelltext",
        "ASM-Zuordnung"]


def schreibe_csv(tab, pfad):
    with open(pfad, "w", newline="", encoding="utf-8-sig") as f:
        w = csv.writer(f, delimiter=";")
        w.writerow(KOPF)
        w.writerows(tab)


MD_TITEL = "Testübersicht"
VERSION_RE = re.compile(r"^Version: (\d+)\.(\d{2})-(\d{3})$", re.M)


def md_einleitung(voll, maxlen):
    """Zwei kurze Absaetze vor der Tabelle: Herkunft und Lesart der Spalten."""
    if voll:
        inhalt = "alle `#DOK`-Zeilen der Testdatei"
    elif maxlen:
        inhalt = (f"die erste `#DOK`-Zeile der Testdatei (höchstens {maxlen} Zeichen; haben mehrere Tests "
                  "denselben Text, sind weitere `#DOK`-Zeilen angehängt)")
    else:
        inhalt = "die erste `#DOK`-Zeile der Testdatei (haben mehrere Tests denselben Text, sind weitere angehängt)"
    return (
        "Übersicht aller Testdateien in `Tests/`, erzeugt von `test_tabelle.py`. Die Spalte „Inhalt des Tests“ "
        f"nennt {inhalt}; fehlt sie, steht die Tastenfolge in eckigen Klammern oder ein Hinweis in runden Klammern.\n\n"
        "Die Spalte „ASM-Zuordnung“ ist leer, wenn die ASM-Datei denselben Namen wie die Testdatei hat; sonst nennt "
        "sie die Art der Zuordnung (über Trommelinhalt, über Namen oder Voraussetzung: Minimalprogramm auf der Trommel)."
    )


def md_version(pfad, rumpf):
    """Version fuer die MD-Datei nach C8205-Versionierung.md (Abschnitt 3): ohne vorhandene Versionszeile
    0.01-001; gleicher Inhalt wie in der vorhandenen Datei: Version bleibt; sonst Build + 1."""
    alt = ""
    if os.path.exists(pfad):
        with open(pfad, encoding="utf-8") as f:
            alt = f.read()
    m = VERSION_RE.search(alt)
    if not m:
        return "0.01-001"
    if alt[m.end():] == "\n\n" + rumpf:
        return f"{m[1]}.{m[2]}-{m[3]}"
    if int(m[3]) >= 999:
        sys.exit("Build 999 erreicht: laut C8205-Versionierung.md vor der naechsten Aenderung den Projektinhaber fragen.")
    return f"{m[1]}.{m[2]}-{int(m[3]) + 1:03d}"


def schreibe_md(tab, pfad, voll=False, maxlen=140):
    def e(s):
        return s.replace("|", "\\|").replace("\n", "<br>")
    zeilen = ["| " + " | ".join(KOPF) + " |", "|" + "---|" * len(KOPF)]
    zeilen += ["| " + " | ".join(e(x) for x in r) + " |" for r in tab]
    rumpf = md_einleitung(voll, maxlen) + "\n\n" + "\n".join(zeilen) + "\n"
    version = md_version(pfad, rumpf)
    with open(pfad, "w", encoding="utf-8") as f:
        f.write(f"# {MD_TITEL}\n\nVersion: {version}\n\n" + rumpf)
    print(f"MD-Version: {version}")


def schreibe_html(tab, pfad):
    with open(pfad, "w", encoding="utf-8") as f:
        f.write("<!DOCTYPE html><meta charset='utf-8'><title>Testübersicht</title>"
                "<style>body{font:14px sans-serif}table{border-collapse:collapse}"
                "td,th{border:1px solid #aaa;padding:4px 8px;vertical-align:top}"
                "th{background:#eee}</style><table><tr>"
                + "".join(f"<th>{html.escape(k)}</th>" for k in KOPF) + "</tr>\n")
        for r in tab:
            f.write("<tr>" + "".join(
                f"<td>{html.escape(x).replace(chr(10), '<br>')}</td>" for x in r) + "</tr>\n")
        f.write("</table>")


def schreibe_xlsx(tab, pfad):
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Alignment, Font
    except ImportError:
        print("openpyxl fehlt -> xlsx übersprungen (pip install openpyxl)", file=sys.stderr)
        return False
    wb = Workbook()
    ws = wb.active
    ws.title = "Tests"
    ws.append(KOPF)
    for r in tab:
        ws.append(list(r))
    for c in ws[1]:
        c.font = Font(bold=True)
    for sp, breite in zip("ABCDE", (42, 70, 60, 50, 30)):
        ws.column_dimensions[sp].width = breite
    for row in ws.iter_rows(min_row=1):
        for c in row:
            c.alignment = Alignment(wrap_text=True, vertical="top")
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    wb.save(pfad)
    return True


def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("--basis", default=None, help="Verzeichnis mit Tests/, Assembler/ ... (Standard: automatisch suchen)")
    ap.add_argument("--format", nargs="+", default=["csv", "md", "xlsx"],
                    choices=["csv", "md", "xlsx", "html"])
    ap.add_argument("--ausgabe", default="test_tabelle", help="Dateiname ohne Endung")
    ap.add_argument("--max-inhalt", type=int, default=140,
                    help="Hoechstlaenge der Spalte 'Inhalt des Tests' in Zeichen (0 = unbegrenzt; Standard 140)")
    ap.add_argument("--voll", action="store_true",
                    help="alle #DOK-Zeilen ungekuerzt (Verhalten vor der Kuerzung)")
    a = ap.parse_args()

    basis = finde_basis(a.basis)
    print("Basis:", basis)
    tab = sammle(basis, a.voll, a.max_inhalt)
    ausgabe = a.ausgabe if os.path.isabs(a.ausgabe) else os.path.join(SKRIPTDIR, a.ausgabe)
    for fmt in a.format:
        ziel = f"{ausgabe}.{fmt}"
        ok = {"csv": schreibe_csv,
              "md": lambda t, z: schreibe_md(t, z, a.voll, a.max_inhalt),
              "html": schreibe_html, "xlsx": schreibe_xlsx}[fmt](tab, ziel)
        if ok is not False:
            print("geschrieben:", ziel)
    berichte_luecken(tab)
    ohne_asm = sum(1 for r in tab if not r[3])
    print(f"{len(tab)} Tests, davon {len(tab) - ohne_asm} mit ASM-Quelltext, {ohne_asm} ohne.")


if __name__ == "__main__":
    main()
